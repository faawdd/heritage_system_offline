# core/views.py
from django.shortcuts import render, redirect
from django.template.loader import render_to_string
from .models import (
    HeritageSite,
    ImmovableHeritage,
    InspectionRecord,
    ProjectAudit,
    Coordinate,
    KmlUploadRecord,
    LandUseProjectApproval,
    LandUseProjectFieldPhoto,
    LandUseProjectOperationLog,
)
from .ovkml_converter import parse_ovkml, build_csv_outputs
from .land_project_services import (
    verify_project_spatial_safety,
    build_project_media_path,
    get_status_controls,
    apply_workflow_action,
)
import base64
import hashlib
import hmac
import json
import csv
from datetime import datetime
from urllib.parse import quote, urlsplit, urlunsplit, parse_qsl, urlencode
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth import get_user_model, login as auth_login
from django.db import OperationalError, ProgrammingError
from django.db.models import Count, Q
from django.http import JsonResponse, HttpResponse, HttpResponseForbidden, FileResponse
from django.utils import timezone
from django.conf import settings
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_POST
from django.core.files.storage import default_storage
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from PIL import Image as PILImage
import os
import io
import uuid
import re
import math
import logging
import zipfile
import xml.etree.ElementTree as ET
from types import SimpleNamespace
from django.shortcuts import get_object_or_404
from django.contrib import messages
from heritage_system.version import VERSION, VERSION_HISTORY
from .permission_decorators import is_management_admin, is_limited_admin
from utils.dem_handler import describe_tile, get_dem_elevation

User = get_user_model()
logger = logging.getLogger(__name__)


LAND_PROJECT_ACTION_LABELS = {
    'create': '收文登记',
    'upload_kml': '上传KML',
    'upload_misc_zip': '上传杂项ZIP',
    'upload_field_photo': '上传现场照片',
    'upload_archaeology_report': '上传考古报告',
    'verify_spatial_safety': '执行空间核验',
    'complete_field_check': '提交现场勘查完成',
    'submit_city_request': '录入县局请示并提交市局',
    'record_city_reply': '录入市局复函',
    'submit_archaeology_request': '发起考古流转',
    'record_archaeology_reply': '录入考古批复结果',
    'archive_case': '办结归档',
}


def _build_payload_doc_nums(payload):
    keys = [
        'shanshan_request_num',
        'city_reply_num',
        'archaeology_request_num',
        'region_approval_num',
        'city_final_reply_num',
        'final_reply_to_company',
    ]
    result = {}
    for key in keys:
        value = payload.get(key)
        if value:
            result[key] = value
    return result


def _record_land_project_operation(project, user, action, payload=None, status_before='', status_after=''):
    payload_data = payload if isinstance(payload, dict) else {}
    LandUseProjectOperationLog.objects.create(
        project=project,
        operator=user if getattr(user, 'is_authenticated', False) else None,
        action=action,
        action_label=LAND_PROJECT_ACTION_LABELS.get(action, action),
        payload=payload_data,
        status_before=status_before or '',
        status_after=status_after or '',
    )


def _b64url_decode(value):
    padding = '=' * (-len(value) % 4)
    return base64.urlsafe_b64decode((value + padding).encode('utf-8'))


def _decode_fastapi_token(token):
    try:
        payload_str, signature = token.split('.', 1)
    except ValueError:
        return None

    expected_signature = hmac.new(
        settings.SECRET_KEY.encode('utf-8'),
        payload_str.encode('utf-8'),
        hashlib.sha256,
    ).hexdigest()
    if not hmac.compare_digest(signature, expected_signature):
        return None

    try:
        payload = json.loads(_b64url_decode(payload_str).decode('utf-8'))
    except Exception:
        return None

    if payload.get('exp', 0) < timezone.now().timestamp():
        return None
    return payload


def mobile_kml_entry_view(request):
    token = (request.GET.get('token') or '').strip()
    target_path = request.GET.get('next', '/admin/kml-overlay-check/')
    if not token:
        return HttpResponseForbidden('缺少登录凭证')

    payload = _decode_fastapi_token(token)
    if not payload:
        return HttpResponseForbidden('登录凭证无效或已过期')

    user = User.objects.filter(id=payload.get('user_id'), is_active=True).first()
    if not user:
        return HttpResponseForbidden('用户不存在或已禁用')

    if not is_management_admin(user):
        return HttpResponseForbidden('当前账号无权使用KML叠加检查')

    auth_login(request, user, backend='django.contrib.auth.backends.ModelBackend')

    split_result = urlsplit(target_path)
    query_pairs = parse_qsl(split_result.query, keep_blank_values=True)
    query_map = dict(query_pairs)

    for key in ('app_loc_lon', 'app_loc_lat', 'app_loc_alt', 'app_loc_accuracy', 'app_loc_ts'):
        value = (request.GET.get(key) or '').strip()
        if value and key not in query_map:
            query_pairs.append((key, value))

    if 'token' not in query_map:
        query_pairs.append(('token', token))

    redirect_url = urlunsplit((
        split_result.scheme,
        split_result.netloc,
        split_result.path,
        urlencode(query_pairs),
        split_result.fragment,
    ))
    return redirect(redirect_url)


def mobile_collect_entry_view(request):
    token = (request.GET.get('token') or '').strip()
    target_path = request.GET.get('next', '/mobile/collect/')
    if not token:
        return HttpResponseForbidden('缺少登录凭证')

    payload = _decode_fastapi_token(token)
    if not payload:
        return HttpResponseForbidden('登录凭证无效或已过期')

    user = User.objects.filter(id=payload.get('user_id'), is_active=True).first()
    if not user:
        return HttpResponseForbidden('用户不存在或已禁用')

    is_admin = user.is_superuser or user.groups.filter(name='管理员').exists() or user.groups.filter(name='超级管理员').exists()
    if not is_admin:
        return HttpResponseForbidden('当前账号无权使用不可移动文物采集管理')

    auth_login(request, user, backend='django.contrib.auth.backends.ModelBackend')
    separator = '&' if '?' in target_path else '?'
    return redirect(f"{target_path}{separator}token={quote(token)}")


@staff_member_required
def system_version_api(request):
    """获取系统版本号信息 API"""
    return JsonResponse({
        'status': 'success',
        'version': VERSION,
        'version_history': VERSION_HISTORY,
    })

TOWNSHIP_NORMALIZATION_RULES = [
    ('东巴扎回族乡', '东巴扎回族乡'),
    ('东巴扎乡', '东巴扎回族乡'),
    ('火车站镇', '火车站镇'),
    ('吐峪沟乡', '吐峪沟乡'),
    ('吐峪沟镇', '吐峪沟乡'),
    ('七克台镇', '七克台镇'),
    ('七克台乡', '七克台镇'),
    ('七台镇', '七克台镇'),
    ('连木沁镇', '连木沁镇'),
    ('连木沁乡', '连木沁镇'),
    ('达朗坎乡', '达朗坎乡'),
    ('达浪坎乡', '达朗坎乡'),
    ('鲁克沁镇', '鲁克沁镇'),
    ('辟展镇', '辟展镇'),
    ('辟展乡', '辟展镇'),
    ('鄯善镇', '鄯善镇'),
    ('迪坎镇', '迪坎镇'),
    ('迪坎乡', '迪坎镇'),
]

TOWNSHIP_STANDARD_TO_KEYWORDS = {}
for keyword, standard_name in TOWNSHIP_NORMALIZATION_RULES:
    TOWNSHIP_STANDARD_TO_KEYWORDS.setdefault(standard_name, set()).add(keyword)


@staff_member_required
def heritage_detail_view(request, pk):
    """旧后台详情页已迁移到 Vue，保留兼容入口。"""
    return redirect(f'/static/frontend/heritage/{pk}')


@staff_member_required
def heritage_boundary_export_view(request, pk):
    """单个不可移动文物边界导出（四普系统）：支持 CSV / KMZ。"""
    heritage = get_object_or_404(HeritageSite, pk=pk)

    if request.method != 'POST':
        return redirect('heritage_detail', pk=pk)

    action = (request.POST.get('action') or '').strip()
    cookie = (request.POST.get('sipu_cookie') or '').strip()
    user_county = (request.POST.get('sipu_county') or '').strip()

    if not cookie:
        messages.error(request, '请先填写四普系统 Cookie，再执行单文物边界导出。')
        return redirect('heritage_detail', pk=pk)

    combined_conflicts = [{
        'feature_source': '单文物导出',
        'site_id': heritage.id,
        'site_name': heritage.name,
        'site_level': heritage.level,
        'site_longitude': heritage.longitude,
        'site_latitude': heritage.latitude,
    }]

    # 构造最小 title 对象以复用统一命名逻辑
    selected_records = [type('ExportRecord', (), {'title': heritage.name})()]

    if action == 'export_single_boundary_csv':
        return _build_boundary_points_csv(combined_conflicts, selected_records, cookie, user_county)

    if action == 'export_single_boundary_kmz':
        return _build_boundary_points_kmz(combined_conflicts, selected_records, cookie, user_county)

    messages.error(request, '未知导出操作。')
    return redirect('heritage_detail', pk=pk)

@staff_member_required
def admin_index_view(request):
    """旧后台首页已迁移到 Vue，保留兼容入口。"""
    return redirect('/static/frontend/dashboard')


def admin_direct_entry_block_view(request):
    """禁止直接进入 Django Admin，统一走 Vue 系统登录后进入。"""
    return redirect('/static/frontend/?redirect=/system/admin')


@staff_member_required  # 确保只有登录后台的人能看
def heritage_map_view(request):
    """旧地图页已迁移到 Vue，保留兼容入口。"""
    return redirect('/static/frontend/heritage/map')


def kml_overlay_check_view(request):
    """旧 KML 管理页已迁移到 Vue，保留兼容入口。"""
    query_string = request.META.get('QUERY_STRING', '')
    target = '/static/frontend/gis/kml-management'
    if query_string:
        target = f'{target}?{query_string}'
    return redirect(target)


def _normalize_threshold(raw_value, default=50, min_value=1, max_value=5000):
    try:
        value = int(raw_value)
    except (TypeError, ValueError):
        return default
    return max(min_value, min(max_value, value))


def _parse_coordinate_text(coord_text):
    points = []
    if not coord_text:
        return points

    for token in str(coord_text).replace('\n', ' ').replace('\t', ' ').split():
        parts = token.split(',')
        if len(parts) < 2:
            continue
        try:
            lon = float(parts[0])
            lat = float(parts[1])
        except (TypeError, ValueError):
            continue
        points.append((lon, lat))
    return points


def _local_tag_name(tag):
    return tag.split('}', 1)[-1] if '}' in tag else tag


def _find_first_coordinates_text(node):
    for child in node.iter():
        if _local_tag_name(child.tag) == 'coordinates' and child.text:
            return child.text
    return ''


def _coord_points_equal(point_a, point_b, tolerance=1e-7):
    if not point_a or not point_b or len(point_a) < 2 or len(point_b) < 2:
        return False
    return abs(point_a[0] - point_b[0]) <= tolerance and abs(point_a[1] - point_b[1]) <= tolerance


def _ensure_closed_ring(points):
    ring_points = list(points or [])
    if len(ring_points) < 3:
        return []
    if not _coord_points_equal(ring_points[0], ring_points[-1]):
        ring_points.append(ring_points[0])
    if len(ring_points) < 4:
        return []
    return ring_points


def _is_closed_ring_points(points):
    if not isinstance(points, list) or len(points) < 4:
        return False
    return _coord_points_equal(points[0], points[-1])


def _is_nearly_closed_ring_points(points, tolerance=1e-5):
    if not isinstance(points, list) or len(points) < 3:
        return False
    return _coord_points_equal(points[0], points[-1], tolerance=tolerance)


def _extract_polygon_rings(polygon_node):
    rings = []

    # 优先按 outer/innerBoundaryIs 提取，保证孔洞顺序正确。
    for boundary_tag in ('outerBoundaryIs', 'innerBoundaryIs'):
        for boundary in polygon_node:
            if _local_tag_name(boundary.tag) != boundary_tag:
                continue
            for ring in boundary:
                if _local_tag_name(ring.tag) != 'LinearRing':
                    continue
                ring_points = _ensure_closed_ring(_parse_coordinate_text(_find_first_coordinates_text(ring)))
                if ring_points:
                    rings.append(ring_points)

    if rings:
        return rings

    for ring in polygon_node.iter():
        if _local_tag_name(ring.tag) != 'LinearRing':
            continue
        ring_points = _ensure_closed_ring(_parse_coordinate_text(_find_first_coordinates_text(ring)))
        if ring_points:
            rings.append(ring_points)
    return rings


def _extract_geometry_features(geom_node, feature_name, source_name):
    geom_type = _local_tag_name(geom_node.tag)
    features = []

    if geom_type == 'Point':
        points = _parse_coordinate_text(_find_first_coordinates_text(geom_node))
        if points:
            features.append({
                'name': feature_name,
                'geometry_type': 'Point',
                'coordinates': points[0],
                'source': source_name,
            })
        return features

    if geom_type == 'LineString':
        points = _parse_coordinate_text(_find_first_coordinates_text(geom_node))
        if not points:
            return features

        if _is_closed_ring_points(points) or _is_nearly_closed_ring_points(points):
            polygon_ring = _ensure_closed_ring(points)
            if not polygon_ring:
                return features
            features.append({
                'name': feature_name,
                'geometry_type': 'Polygon',
                'coordinates': [polygon_ring],
                'source': source_name,
            })
        else:
            features.append({
                'name': feature_name,
                'geometry_type': 'LineString',
                'coordinates': points,
                'source': source_name,
            })
        return features

    if geom_type == 'LinearRing':
        ring_points = _ensure_closed_ring(_parse_coordinate_text(_find_first_coordinates_text(geom_node)))
        if ring_points:
            features.append({
                'name': feature_name,
                'geometry_type': 'Polygon',
                'coordinates': [ring_points],
                'source': source_name,
            })
        return features

    if geom_type == 'Polygon':
        rings = _extract_polygon_rings(geom_node)
        if rings:
            features.append({
                'name': feature_name,
                'geometry_type': 'Polygon',
                'coordinates': rings,
                'source': source_name,
            })
        return features

    if geom_type in {'MultiGeometry', 'GeometryCollection'}:
        for child_geom in geom_node:
            child_type = _local_tag_name(child_geom.tag)
            if child_type in {'Point', 'LineString', 'LinearRing', 'Polygon', 'MultiGeometry', 'GeometryCollection'}:
                features.extend(_extract_geometry_features(child_geom, feature_name, source_name))
        return features

    return features


def _extract_features_from_kml_xml(xml_text, source_name):
    try:
        root = ET.fromstring(xml_text)
    except ET.ParseError:
        return []

    features = []
    for placemark in root.iter():
        if _local_tag_name(placemark.tag) != 'Placemark':
            continue

        feature_name = 'KML要素'
        for child in placemark:
            if _local_tag_name(child.tag) == 'name' and child.text and child.text.strip():
                feature_name = child.text.strip()
                break

        for child in placemark:
            child_type = _local_tag_name(child.tag)
            if child_type in {'Point', 'LineString', 'LinearRing', 'Polygon', 'MultiGeometry', 'GeometryCollection'}:
                features.extend(_extract_geometry_features(child, feature_name, source_name))

    return features


def _extract_features_from_upload(filename, content_bytes):
    lower_name = (filename or '').lower()
    features = []

    if lower_name.endswith('.kmz') or lower_name.endswith('.ovkmz'):
        try:
            with zipfile.ZipFile(io.BytesIO(content_bytes), 'r') as zf:
                for member in zf.namelist():
                    if not member.lower().endswith('.kml'):
                        continue
                    try:
                        member_bytes = zf.read(member)
                    except Exception:
                        # 单个成员损坏时跳过，避免整包失败。
                        continue
                    features.extend(_extract_features_from_kml_xml(member_bytes, f"{filename}:{member}"))
        except (zipfile.BadZipFile, RuntimeError, OSError):
            # 兼容部分浏览器/端将 KML 误命名为 KMZ 的情况，回退按 KML 文本解析。
            features.extend(_extract_features_from_kml_xml(content_bytes, filename))
    else:
        features.extend(_extract_features_from_kml_xml(content_bytes, filename))

    return features


def _haversine_m(lat1, lon1, lat2, lon2):
    radius = 6371000
    lat1_rad = math.radians(lat1)
    lat2_rad = math.radians(lat2)
    dlat = math.radians(lat2 - lat1)
    dlon = math.radians(lon2 - lon1)
    a = math.sin(dlat / 2) ** 2 + math.cos(lat1_rad) * math.cos(lat2_rad) * math.sin(dlon / 2) ** 2
    return 2 * radius * math.atan2(math.sqrt(a), math.sqrt(1 - a))


def _distance_point_to_segment_m(lat, lon, lat1, lon1, lat2, lon2):
    meters_per_deg_lat = 111320
    meters_per_deg_lon = 111320 * math.cos(math.radians((lat1 + lat2) / 2))

    px = lon * meters_per_deg_lon
    py = lat * meters_per_deg_lat
    x1 = lon1 * meters_per_deg_lon
    y1 = lat1 * meters_per_deg_lat
    x2 = lon2 * meters_per_deg_lon
    y2 = lat2 * meters_per_deg_lat

    dx = x2 - x1
    dy = y2 - y1
    if dx == 0 and dy == 0:
        return math.hypot(px - x1, py - y1)

    t = ((px - x1) * dx + (py - y1) * dy) / (dx * dx + dy * dy)
    t = max(0.0, min(1.0, t))
    proj_x = x1 + t * dx
    proj_y = y1 + t * dy
    return math.hypot(px - proj_x, py - proj_y)


def _distance_to_linestring_m(site_lon, site_lat, line_coords):
    if not line_coords:
        return float('inf')
    if len(line_coords) == 1:
        lon, lat = line_coords[0]
        return _haversine_m(site_lat, site_lon, lat, lon)

    min_distance = float('inf')
    for idx in range(len(line_coords) - 1):
        lon1, lat1 = line_coords[idx]
        lon2, lat2 = line_coords[idx + 1]
        distance = _distance_point_to_segment_m(site_lat, site_lon, lat1, lon1, lat2, lon2)
        min_distance = min(min_distance, distance)
    return min_distance


def _distance_to_polygon_boundary_m(site_lon, site_lat, polygon_rings):
    """计算点到多边形边界（外环+内环）的最短距离（米）。"""
    if not polygon_rings:
        return float('inf')

    min_distance = float('inf')
    for ring in polygon_rings:
        if not ring:
            continue

        ring_points = list(ring)
        # KML 线环可能未闭合，统一补齐闭合段
        if len(ring_points) >= 2 and ring_points[0] != ring_points[-1]:
            ring_points.append(ring_points[0])

        ring_distance = _distance_to_linestring_m(site_lon, site_lat, ring_points)
        min_distance = min(min_distance, ring_distance)

    return min_distance


def _distance_to_multipolygon_boundary_m(site_lon, site_lat, multi_polygon_coords):
    if not multi_polygon_coords:
        return float('inf')

    min_distance = float('inf')
    for polygon_rings in multi_polygon_coords:
        min_distance = min(min_distance, _distance_to_polygon_boundary_m(site_lon, site_lat, polygon_rings))
    return min_distance


def _is_point_in_ring(lon, lat, ring):
    if not ring or len(ring) < 3:
        return False

    inside = False
    j = len(ring) - 1
    for i in range(len(ring)):
        xi, yi = ring[i]
        xj, yj = ring[j]
        if ((yi > lat) != (yj > lat)) and (lon < (xj - xi) * (lat - yi) / ((yj - yi) or 1e-12) + xi):
            inside = not inside
        j = i
    return inside


def _is_point_in_polygon(lon, lat, polygon_rings):
    if not polygon_rings:
        return False
    if not _is_point_in_ring(lon, lat, polygon_rings[0]):
        return False
    for hole_ring in polygon_rings[1:]:
        if _is_point_in_ring(lon, lat, hole_ring):
            return False
    return True


def _load_conflict_site_points():
    rows = HeritageSite.objects.exclude(longitude__isnull=True).exclude(latitude__isnull=True).values(
        'id', 'name', 'level', 'longitude', 'latitude'
    )
    site_points = []
    for row in rows:
        try:
            site_points.append(
                {
                    'id': row['id'],
                    'name': row['name'],
                    'level': row['level'],
                    'longitude': float(row['longitude']),
                    'latitude': float(row['latitude']),
                }
            )
        except (TypeError, ValueError):
            continue
    return site_points


def _analyze_conflicts(features, threshold_m, site_points=None):
    if site_points is None:
        site_points = _load_conflict_site_points()

    threshold = float(threshold_m)
    conflicts = []

    for feature in features:
        feature_type = feature.get('geometry_type')
        feature_name = feature.get('name') or 'KML要素'
        feature_source = feature.get('source') or ''
        coords = feature.get('coordinates')

        for site in site_points:
            site_lon = site['longitude']
            site_lat = site['latitude']

            matched = False
            relation = ''
            distance_m = None

            if feature_type == 'Point' and isinstance(coords, (list, tuple)) and len(coords) >= 2:
                lon, lat = float(coords[0]), float(coords[1])
                distance_m = _haversine_m(site_lat, site_lon, lat, lon)
                matched = distance_m <= threshold
                relation = '点距离'
            elif feature_type == 'LineString':
                distance_m = _distance_to_linestring_m(site_lon, site_lat, coords or [])
                matched = distance_m <= threshold
                relation = '线最短距离'
            elif feature_type == 'MultiLineString':
                min_distance = float('inf')
                for line_coords in (coords or []):
                    min_distance = min(min_distance, _distance_to_linestring_m(site_lon, site_lat, line_coords))
                distance_m = min_distance
                matched = distance_m <= threshold
                relation = '线最短距离'
            elif feature_type == 'Polygon':
                inside = _is_point_in_polygon(site_lon, site_lat, coords or [])
                boundary_distance = _distance_to_polygon_boundary_m(site_lon, site_lat, coords or [])
                distance_m = boundary_distance
                matched = inside or (math.isfinite(boundary_distance) and boundary_distance <= threshold)
                relation = '面内包含' if inside else '面边界最短距离'
            elif feature_type == 'MultiPolygon':
                inside = any(_is_point_in_polygon(site_lon, site_lat, polygon) for polygon in (coords or []))
                boundary_distance = _distance_to_multipolygon_boundary_m(site_lon, site_lat, coords or [])
                distance_m = boundary_distance
                matched = inside or (math.isfinite(boundary_distance) and boundary_distance <= threshold)
                relation = '面内包含' if inside else '面边界最短距离'

            if matched:
                conflicts.append({
                    'feature_name': feature_name,
                    'feature_type': feature_type,
                    'feature_source': feature_source,
                    'site_id': site['id'],
                    'site_name': site['name'],
                    'site_level': site['level'],
                    'site_longitude': round(site_lon, 8),
                    'site_latitude': round(site_lat, 8),
                    'relation': relation,
                    'distance_m': None if distance_m is None or not math.isfinite(distance_m) else round(distance_m, 2),
                })

    return conflicts


def _build_conflict_report_csv(conflicts, threshold_m, records=None):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['阈值(米)', threshold_m])
    writer.writerow([])
    writer.writerow(['文件来源', '要素名称', '要素类型', '文物ID', '文物名称', '文物级别', '文物经度', '文物纬度', '冲突关系', '距离(米)'])

    if conflicts:
        for row in conflicts:
            writer.writerow([
                row.get('feature_source', ''),
                row.get('feature_name', ''),
                row.get('feature_type', ''),
                row.get('site_id', ''),
                row.get('site_name', ''),
                row.get('site_level', ''),
                row.get('site_longitude', ''),
                row.get('site_latitude', ''),
                row.get('relation', ''),
                '' if row.get('distance_m') is None else row.get('distance_m'),
            ])
    else:
        writer.writerow(['-', '-', '-', '-', '无冲突', '-', '-', '-', '-', '-'])

    response = HttpResponse(output.getvalue(), content_type='text/csv; charset=utf-8-sig')
    date_str = timezone.now().strftime('%Y%m%d')
    if records and len(records) == 1:
        first_title = records[0].title
        report_name = f'{date_str}{first_title}查询报告'
    elif records and len(records) > 1:
        first_title = records[0].title
        report_name = f'{date_str}{first_title}等查询报告'
    else:
        report_name = f'kml_conflict_report_{timezone.now().strftime("%Y%m%d_%H%M%S")}'
    # RFC 5987 UTF-8 编码处理中文文件名
    from urllib.parse import quote
    encoded_name = quote(report_name + '.csv', safe='')
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{encoded_name}"
    return response


def _build_conflict_sites_kml(conflicts, threshold_m, records=None):
    """将冲突文物点导出为奥维互动地图兼容的 KML 格式（按 site_id 去重，仅导出唯一文物点）。"""
    seen_ids = set()
    unique_sites = []
    for row in conflicts:
        sid = row.get('site_id')
        if sid and sid not in seen_ids:
            seen_ids.add(sid)
            unique_sites.append(row)

    date_str = timezone.now().strftime('%Y%m%d')
    if records and len(records) == 1:
        kml_name = f'{date_str}{records[0].title}冲突文物点'
    elif records and len(records) > 1:
        kml_name = f'{date_str}{records[0].title}等冲突文物点'
    else:
        kml_name = f'冲突文物点_{timezone.now().strftime("%Y%m%d_%H%M%S")}'

    ET.register_namespace('', 'http://www.opengis.net/kml/2.2')
    ET.register_namespace('gx', 'http://www.google.com/kml/ext/2.2')
    ns = 'http://www.opengis.net/kml/2.2'

    kml_root = ET.Element(f'{{{ns}}}kml')
    doc = ET.SubElement(kml_root, f'{{{ns}}}Document')
    ET.SubElement(doc, f'{{{ns}}}name').text = kml_name

    # 奥维红色图标样式
    style = ET.SubElement(doc, f'{{{ns}}}Style')
    style.set('id', 'heritageConflict')
    icon_style = ET.SubElement(style, f'{{{ns}}}IconStyle')
    ET.SubElement(icon_style, f'{{{ns}}}color').text = 'ff0000ff'  # ABGR 红色
    ET.SubElement(icon_style, f'{{{ns}}}scale').text = '1.2'
    icon_el = ET.SubElement(icon_style, f'{{{ns}}}Icon')
    ET.SubElement(icon_el, f'{{{ns}}}href').text = 'http://maps.google.com/mapfiles/kml/paddle/red-circle.png'
    label_style = ET.SubElement(style, f'{{{ns}}}LabelStyle')
    ET.SubElement(label_style, f'{{{ns}}}color').text = 'ff0000ff'
    ET.SubElement(label_style, f'{{{ns}}}scale').text = '0.9'

    for site in unique_sites:
        lon = site.get('site_longitude')
        lat = site.get('site_latitude')
        if lon is None or lat is None:
            continue
        pm = ET.SubElement(doc, f'{{{ns}}}Placemark')
        ET.SubElement(pm, f'{{{ns}}}name').text = str(site.get('site_name', '未命名文物'))
        ET.SubElement(pm, f'{{{ns}}}description').text = (
            f'文物级别: {site.get("site_level", "")}\n'
            f'文物ID: {site.get("site_id", "")}\n'
            f'阈值: {threshold_m}米'
        )
        ET.SubElement(pm, f'{{{ns}}}styleUrl').text = '#heritageConflict'
        point_el = ET.SubElement(pm, f'{{{ns}}}Point')
        ET.SubElement(point_el, f'{{{ns}}}coordinates').text = f'{lon},{lat},0'

    kml_bytes = ET.tostring(kml_root, encoding='unicode', xml_declaration=False)
    kml_content = '<?xml version="1.0" encoding="UTF-8"?>\n' + kml_bytes

    response = HttpResponse(kml_content, content_type='application/vnd.google-earth.kml+xml; charset=utf-8')
    encoded_name = quote(kml_name + '.kml', safe='')
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{encoded_name}"
    return response


def _reanalyze_kml_records(records, threshold, use_cache=True):
    combined_conflicts = []
    updated_count = 0
    failed_count = 0
    failed_items = []
    records_to_update = []
    site_points = _load_conflict_site_points()

    for record in records:
        try:
            if use_cache and int(record.threshold_m or 0) == int(threshold) and record.report_json:
                cached_payload = json.loads(record.report_json)
                cached_conflicts = cached_payload.get('conflicts', [])
                if isinstance(cached_conflicts, list):
                    combined_conflicts.extend(cached_conflicts)
                    updated_count += 1
                    continue

            if not record.source_file or not record.source_file.name:
                raise FileNotFoundError('记录未绑定源文件')
            try:
                file_exists = record.source_file.storage.exists(record.source_file.name)
            except Exception:
                file_exists = True
            if not file_exists:
                raise FileNotFoundError(f'源文件不存在: {record.source_file.name}')

            with record.source_file.open('rb') as source:
                content = source.read()
            file_name_for_parse = record.source_file.name or record.title or ''
            features = _extract_features_from_upload(file_name_for_parse, content)
            conflicts = _analyze_conflicts(features, threshold, site_points=site_points)
        except Exception as exc:
            logger.exception('KML重分析失败: record_id=%s title=%s source=%s', record.id, record.title, getattr(record.source_file, 'name', ''))
            failed_count += 1
            failed_items.append({
                'title': record.title or f'记录#{record.id}',
                'error': str(exc) or exc.__class__.__name__,
            })
            continue

        record.threshold_m = threshold
        record.feature_count = len(features)
        record.conflict_count = len(conflicts)
        record.report_json = json.dumps({
            'threshold_m': threshold,
            'feature_count': len(features),
            'conflict_count': len(conflicts),
            'generated_at': timezone.now().isoformat(),
            'conflicts': conflicts,
        }, ensure_ascii=False)
        records_to_update.append(record)

        updated_count += 1
        combined_conflicts.extend(conflicts)

    if records_to_update:
        KmlUploadRecord.objects.bulk_update(
            records_to_update,
            ['threshold_m', 'feature_count', 'conflict_count', 'report_json', 'updated_at'],
        )

    return combined_conflicts, updated_count, failed_count, failed_items


def _is_admin_user(user):
    return user.is_authenticated and is_management_admin(user)


# ──────────────────────────────────────────────────────────────────────────────
# 四普系统边界坐标导出辅助函数
# ──────────────────────────────────────────────────────────────────────────────

_SIPU_HOST = '202.41.243.152:9046'
_SIPU_BASE = f'http://{_SIPU_HOST}'


def _sipu_search_culrid(site_name: str, cookie: str, user_county: str = '') -> list:
    """通过四普系统搜索接口，根据文物名称获取候选记录列表（含 id/culRid）。"""
    import urllib.request
    import urllib.parse

    url = f'{_SIPU_BASE}/immovableListController.do?queryRelicList'
    form_data = {
        'page': '1',
        'pageSize': '5',
        'searchInputValue': site_name,
        'sortField': 'update_date',
        'sortType': 'desc',
        'backStatus': '0',
    }
    if user_county:
        form_data['userCounty'] = user_county

    body = urllib.parse.urlencode(form_data).encode('utf-8')
    req = urllib.request.Request(
        url,
        data=body,
        method='POST',
        headers={
            'Content-Type': 'application/x-www-form-urlencoded; charset=UTF-8',
            'X-Requested-With': 'XMLHttpRequest',
            'Accept': 'application/json, text/javascript, */*; q=0.01',
            'Cookie': cookie,
            'Host': _SIPU_HOST,
            'Origin': _SIPU_BASE,
            'Referer': f'{_SIPU_BASE}/immovableListController.do?immovableList',
            'User-Agent': 'Mozilla/5.0 (compatible; HeritageSystem/1.0)',
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=15) as resp:
            raw = resp.read()
        data = json.loads(raw.decode('utf-8', errors='replace'))
        return data.get('data') or []
    except Exception:
        return []


def _sipu_fetch_boundary_points(cul_rid: str, cookie: str) -> list:
    """通过四普系统坐标列表接口，获取指定文物的边界点（measurePointType=1）。
    由于接口不支持按 measurePointType 过滤，在客户端过滤。
    """
    import urllib.request
    import urllib.parse

    all_rows = []
    page = 1
    page_count = 50  # 一次多取，减少请求次数

    while True:
        url = (
            f'{_SIPU_BASE}/tBBdataPointsController.do'
            f'?getData&currpage={page}&pagecount={page_count}'
        )
        body = urllib.parse.urlencode({'culRid': cul_rid}).encode('utf-8')
        req = urllib.request.Request(
            url,
            data=body,
            method='POST',
            headers={
                'Content-Type': 'application/x-www-form-urlencoded; charset=UTF-8',
                'X-Requested-With': 'XMLHttpRequest',
                'Accept': 'application/json, text/javascript, */*; q=0.01',
                'Cookie': cookie,
                'Host': _SIPU_HOST,
                'Origin': _SIPU_BASE,
                'Referer': (
                    f'{_SIPU_BASE}/tBBdataBasicController.do'
                    f'?tBBdataPointsView&type=&culRid={urllib.parse.quote(cul_rid)}'
                ),
                'User-Agent': 'Mozilla/5.0 (compatible; HeritageSystem/1.0)',
            },
        )
        try:
            with urllib.request.urlopen(req, timeout=15) as resp:
                raw = resp.read()
            data = json.loads(raw.decode('utf-8', errors='replace'))
        except Exception:
            break

        rows = data.get('rows') or []
        total = int(data.get('total') or 0)
        all_rows.extend(rows)

        if len(all_rows) >= total or not rows:
            break
        page += 1

    # 过滤边界点：优先 measurePointType == "1"；若无则兼容部分数据使用的 "9"
    boundary = [r for r in all_rows if str(r.get('measurePointType', '')) == '1']
    if not boundary:
        boundary = [r for r in all_rows if str(r.get('measurePointType', '')) == '9']
    return boundary


def _build_boundary_points_csv(combined_conflicts, selected_records, cookie: str, user_county: str = '') -> HttpResponse:
    """
    对冲突文物点按文件分组，逐个调用四普系统接口获取边界坐标，
    导出为一张 CSV 表格（含文件分组列）。
    """
    # 按来源 KML 文件聚合冲突文物（site_id 去重）
    from collections import OrderedDict

    # 建立 {feature_source: [site_id, ...]} 映射（保序、去重）
    source_sites: dict = OrderedDict()
    site_meta: dict = {}  # site_id -> {name, level, longitude, latitude}

    for row in combined_conflicts:
        src = row.get('feature_source') or '未知来源'
        sid = row.get('site_id')
        if not sid:
            continue
        source_sites.setdefault(src, [])
        if sid not in source_sites[src]:
            source_sites[src].append(sid)
        if sid not in site_meta:
            site_meta[sid] = {
                'name': row.get('site_name', ''),
                'level': row.get('site_level', ''),
                'longitude': row.get('site_longitude', ''),
                'latitude': row.get('site_latitude', ''),
            }

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        '来源KML文件', '文物名称', '文物级别', '四普文物名称',
        '序号', '点描述', '备注',
        '纬度(十进制)', '经度(十进制)', '海拔',
        '纬度度', '纬度分', '纬度秒',
        '经度度', '经度分', '经度秒',
        '出界标记', '距出界距离(m)',
    ])

    # 用于缓存 site_id -> culRid 映射，避免重复搜索
    cul_rid_cache: dict = {}

    for src, site_ids in source_sites.items():
        for sid in site_ids:
            meta = site_meta[sid]
            site_name = meta['name']

            # 查找 culRid
            if sid in cul_rid_cache:
                cul_rid, sipu_name = cul_rid_cache[sid]
            else:
                candidates = _sipu_search_culrid(site_name, cookie, user_county)
                # 精确匹配文物名称；若无精确匹配则取第一条
                matched = next((c for c in candidates if c.get('name') == site_name), None)
                if matched is None and candidates:
                    matched = candidates[0]
                if matched:
                    cul_rid = matched.get('id') or ''
                    sipu_name = matched.get('name') or ''
                else:
                    cul_rid = ''
                    sipu_name = ''
                cul_rid_cache[sid] = (cul_rid, sipu_name)

            if not cul_rid:
                writer.writerow([
                    src, site_name, meta['level'], '（四普系统未找到该文物）',
                    '', '', '', '', '', '', '', '', '', '', '', '', '', '',
                ])
                continue

            points = _sipu_fetch_boundary_points(cul_rid, cookie)
            if not points:
                writer.writerow([
                    src, site_name, meta['level'], sipu_name,
                    '', '', '', '', '', '', '', '', '', '', '', '', '（无边界点数据）', '',
                ])
                continue

            for pt in points:
                writer.writerow([
                    src,
                    site_name,
                    meta['level'],
                    sipu_name,
                    pt.get('counter', ''),
                    pt.get('pointDesc', ''),
                    pt.get('remark', ''),
                    pt.get('lat', ''),
                    pt.get('lng', ''),
                    pt.get('altitude', ''),
                    pt.get('latitude1', ''),
                    pt.get('latitude2', ''),
                    pt.get('latitude3', ''),
                    pt.get('longitude1', ''),
                    pt.get('longitude2', ''),
                    pt.get('longitude3', ''),
                    pt.get('outBody', ''),
                    pt.get('distanceOut', ''),
                ])

    date_str = timezone.now().strftime('%Y%m%d')
    if selected_records and len(selected_records) == 1:
        report_name = f'{date_str}{selected_records[0].title}冲突文物边界坐标'
    elif selected_records:
        report_name = f'{date_str}{selected_records[0].title}等冲突文物边界坐标'
    else:
        report_name = f'冲突文物边界坐标_{timezone.now().strftime("%Y%m%d_%H%M%S")}'

    response = HttpResponse(output.getvalue(), content_type='text/csv; charset=utf-8-sig')
    encoded_name = quote(report_name + '.csv', safe='')
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{encoded_name}"
    return response


def _build_boundary_points_kmz(combined_conflicts, selected_records, cookie: str, user_county: str = '') -> HttpResponse:
    """
    将冲突文物点在四普系统中的边界点（measurePointType=1）导出为 KMZ。
    每个文物点按边界点顺序闭合成面，并以文物名称命名 Placemark。
    """
    from collections import OrderedDict

    source_sites: dict = OrderedDict()
    site_meta: dict = {}

    for row in combined_conflicts:
        src = row.get('feature_source') or '未知来源'
        sid = row.get('site_id')
        if not sid:
            continue
        source_sites.setdefault(src, [])
        if sid not in source_sites[src]:
            source_sites[src].append(sid)
        if sid not in site_meta:
            site_meta[sid] = {
                'name': row.get('site_name', ''),
                'level': row.get('site_level', ''),
            }

    def _point_order_key(item):
        for key in ('snNuM', 'counter'):
            value = item.get(key)
            try:
                return int(value)
            except (TypeError, ValueError):
                continue
        return 0

    def _to_lonlat(item):
        try:
            lon = float(item.get('lng'))
            lat = float(item.get('lat'))
        except (TypeError, ValueError):
            return None
        return lon, lat

    cul_rid_cache: dict = {}
    polygons = []

    for src, site_ids in source_sites.items():
        for sid in site_ids:
            meta = site_meta[sid]
            site_name = meta.get('name') or '未命名文物'

            if sid in cul_rid_cache:
                cul_rid, sipu_name = cul_rid_cache[sid]
            else:
                candidates = _sipu_search_culrid(site_name, cookie, user_county)
                matched = next((c for c in candidates if c.get('name') == site_name), None)
                if matched is None and candidates:
                    matched = candidates[0]
                if matched:
                    cul_rid = matched.get('id') or ''
                    sipu_name = matched.get('name') or site_name
                else:
                    cul_rid = ''
                    sipu_name = site_name
                cul_rid_cache[sid] = (cul_rid, sipu_name)

            if not cul_rid:
                continue

            points = _sipu_fetch_boundary_points(cul_rid, cookie)
            if not points:
                continue

            # 按 groupLink 分区，避免多块墓地被错误串接成一个面
            grouped_points = {}
            for item in points:
                group_key = str(item.get('groupLink') or '1')
                grouped_points.setdefault(group_key, []).append(item)

            rings = []
            for group_key, group_items in grouped_points.items():
                ordered = sorted(group_items, key=_point_order_key)
                ring = []
                for item in ordered:
                    lonlat = _to_lonlat(item)
                    if lonlat is None:
                        continue
                    ring.append(lonlat)

                # 多边形至少需要3个点
                if len(ring) < 3:
                    continue

                # 闭合线环
                if ring[0] != ring[-1]:
                    ring.append(ring[0])

                rings.append({
                    'group_key': group_key,
                    'coords': ring,
                })

            if not rings:
                continue

            polygons.append({
                'name': sipu_name or site_name,
                'source': src,
                'site_level': meta.get('level', ''),
                'cul_rid': cul_rid,
                'rings': rings,
            })

    if not polygons:
        # 无可导出多边形时，返回空 KML 文档，避免下载报错
        polygons = []

    date_str = timezone.now().strftime('%Y%m%d')
    if selected_records and len(selected_records) == 1:
        kmz_name = f'{date_str}{selected_records[0].title}冲突文物边界面'
    elif selected_records:
        kmz_name = f'{date_str}{selected_records[0].title}等冲突文物边界面'
    else:
        kmz_name = f'冲突文物边界面_{timezone.now().strftime("%Y%m%d_%H%M%S")}'

    ET.register_namespace('', 'http://www.opengis.net/kml/2.2')
    ns = 'http://www.opengis.net/kml/2.2'
    kml_root = ET.Element(f'{{{ns}}}kml')
    doc = ET.SubElement(kml_root, f'{{{ns}}}Document')
    ET.SubElement(doc, f'{{{ns}}}name').text = kmz_name

    # 奥维可读的面样式（红边半透明填充）
    style = ET.SubElement(doc, f'{{{ns}}}Style')
    style.set('id', 'conflictBoundaryPolygon')
    line_style = ET.SubElement(style, f'{{{ns}}}LineStyle')
    ET.SubElement(line_style, f'{{{ns}}}color').text = 'ff0000ff'
    ET.SubElement(line_style, f'{{{ns}}}width').text = '2'
    poly_style = ET.SubElement(style, f'{{{ns}}}PolyStyle')
    ET.SubElement(poly_style, f'{{{ns}}}color').text = '4d0000ff'

    for item in polygons:
        pm = ET.SubElement(doc, f'{{{ns}}}Placemark')
        ET.SubElement(pm, f'{{{ns}}}name').text = item['name']
        ET.SubElement(pm, f'{{{ns}}}styleUrl').text = '#conflictBoundaryPolygon'
        total_points = sum(max(len(r['coords']) - 1, 0) for r in item['rings'])
        ET.SubElement(pm, f'{{{ns}}}description').text = (
            f"来源KML: {item['source']}\n"
            f"文物级别: {item['site_level']}\n"
            f"区块数: {len(item['rings'])}\n"
            f"边界点数: {total_points}"
        )

        if len(item['rings']) == 1:
            ring_coords = item['rings'][0]['coords']
            polygon = ET.SubElement(pm, f'{{{ns}}}Polygon')
            ET.SubElement(polygon, f'{{{ns}}}tessellate').text = '1'
            outer = ET.SubElement(polygon, f'{{{ns}}}outerBoundaryIs')
            ring = ET.SubElement(outer, f'{{{ns}}}LinearRing')
            coord_text = ' '.join(f'{lon:.10f},{lat:.10f},0' for lon, lat in ring_coords)
            ET.SubElement(ring, f'{{{ns}}}coordinates').text = coord_text
        else:
            multi_geometry = ET.SubElement(pm, f'{{{ns}}}MultiGeometry')
            for ring_item in item['rings']:
                polygon = ET.SubElement(multi_geometry, f'{{{ns}}}Polygon')
                ET.SubElement(polygon, f'{{{ns}}}tessellate').text = '1'
                outer = ET.SubElement(polygon, f'{{{ns}}}outerBoundaryIs')
                ring = ET.SubElement(outer, f'{{{ns}}}LinearRing')
                coord_text = ' '.join(f'{lon:.10f},{lat:.10f},0' for lon, lat in ring_item['coords'])
                ET.SubElement(ring, f'{{{ns}}}coordinates').text = coord_text

    kml_bytes = ET.tostring(kml_root, encoding='utf-8', xml_declaration=True)

    kmz_buffer = io.BytesIO()
    with zipfile.ZipFile(kmz_buffer, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr('doc.kml', kml_bytes)

    response = HttpResponse(kmz_buffer.getvalue(), content_type='application/vnd.google-earth.kmz')
    encoded_name = quote(kmz_name + '.kmz', safe='')
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{encoded_name}"
    return response


@staff_member_required
def kml_management_view(request):
    """旧 KML 文件管理页已迁移到 Vue，保留兼容入口。"""
    query_string = request.META.get('QUERY_STRING', '')
    target = '/static/frontend/gis/kml-management'
    if query_string:
        target = f'{target}?{query_string}'
    return redirect(target)


def _is_kml_family_filename(filename: str) -> bool:
    lower_name = (filename or '').lower()
    return lower_name.endswith('.kml') or lower_name.endswith('.kmz') or lower_name.endswith('.ovkml') or lower_name.endswith('.ovkmz')


def _build_kml_table_rows(records, output_mode: str):
    rows = []
    for idx, item in enumerate(records, start=1):
        row = {
            'index': idx,
            'project_name': item.project_name,
            'geometry_type': item.geometry_type,
            'vertex_count': item.vertex_count,
            'source_folder': item.source_folder or '-',
            'project_coordinates': item.project_coordinates,
        }
        if output_mode == 'cgcs2000_proj':
            row['coord_a'] = '' if item.cgcs2000_x is None else f'{item.cgcs2000_x:.3f}'
            row['coord_b'] = '' if item.cgcs2000_y is None else f'{item.cgcs2000_y:.3f}'
        else:
            row['coord_a'] = '' if item.target_lon is None else f'{item.target_lon:.10f}'
            row['coord_b'] = '' if item.target_lat is None else f'{item.target_lat:.10f}'
        rows.append(row)
    return rows


def _build_kml_table_csv(rows, output_mode: str) -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    if output_mode == 'cgcs2000_proj':
        writer.writerow(['序号', '要素名称', '几何类型', '顶点数', 'CGCS2000_X(米)', 'CGCS2000_Y(米)', '来源文件夹', '坐标串'])
    else:
        writer.writerow(['序号', '要素名称', '几何类型', '顶点数', '经度', '纬度', '来源文件夹', '坐标串'])

    for row in rows:
        writer.writerow([
            row.get('index', ''),
            row.get('project_name', ''),
            row.get('geometry_type', ''),
            row.get('vertex_count', ''),
            row.get('coord_a', ''),
            row.get('coord_b', ''),
            row.get('source_folder', ''),
            row.get('project_coordinates', ''),
        ])

    return output.getvalue()


def _convert_dxf_bytes_to_kml(dxf_bytes: bytes, doc_name: str):
    try:
        import ezdxf
    except Exception:
        raise RuntimeError('DXF 转换依赖 ezdxf，请先安装：pip install ezdxf')

    import tempfile

    temp_path = ''
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix='.dxf') as tmp:
            tmp.write(dxf_bytes)
            temp_path = tmp.name

        doc = ezdxf.readfile(temp_path)
        msp = doc.modelspace()

        ET.register_namespace('', 'http://www.opengis.net/kml/2.2')
        ns = 'http://www.opengis.net/kml/2.2'
        kml_root = ET.Element(f'{{{ns}}}kml')
        doc_el = ET.SubElement(kml_root, f'{{{ns}}}Document')
        ET.SubElement(doc_el, f'{{{ns}}}name').text = doc_name

        stats = {
            'point_count': 0,
            'line_count': 0,
            'polygon_count': 0,
            'unsupported_count': 0,
        }

        def _add_placemark(name, geometry_type, coords):
            if not coords:
                return
            pm = ET.SubElement(doc_el, f'{{{ns}}}Placemark')
            ET.SubElement(pm, f'{{{ns}}}name').text = name

            if geometry_type == 'Point':
                point = ET.SubElement(pm, f'{{{ns}}}Point')
                x, y = coords[0]
                ET.SubElement(point, f'{{{ns}}}coordinates').text = f'{x:.10f},{y:.10f},0'
                stats['point_count'] += 1
                return

            if geometry_type == 'Polygon':
                polygon = ET.SubElement(pm, f'{{{ns}}}Polygon')
                outer = ET.SubElement(polygon, f'{{{ns}}}outerBoundaryIs')
                ring = ET.SubElement(outer, f'{{{ns}}}LinearRing')
                text = ' '.join(f'{x:.10f},{y:.10f},0' for x, y in coords)
                ET.SubElement(ring, f'{{{ns}}}coordinates').text = text
                stats['polygon_count'] += 1
                return

            line = ET.SubElement(pm, f'{{{ns}}}LineString')
            text = ' '.join(f'{x:.10f},{y:.10f},0' for x, y in coords)
            ET.SubElement(line, f'{{{ns}}}coordinates').text = text
            stats['line_count'] += 1

        for entity in msp:
            etype = entity.dxftype()
            layer_name = getattr(entity.dxf, 'layer', '') or 'DXF'

            if etype == 'POINT':
                location = entity.dxf.location
                _add_placemark(f'{layer_name}-POINT', 'Point', [(float(location.x), float(location.y))])
                continue

            if etype == 'LINE':
                start = entity.dxf.start
                end = entity.dxf.end
                _add_placemark(
                    f'{layer_name}-LINE',
                    'LineString',
                    [(float(start.x), float(start.y)), (float(end.x), float(end.y))],
                )
                continue

            if etype == 'LWPOLYLINE':
                pts = [(float(p[0]), float(p[1])) for p in entity.get_points('xy')]
                if len(pts) < 2:
                    continue
                if bool(entity.closed) and len(pts) >= 3:
                    if pts[0] != pts[-1]:
                        pts.append(pts[0])
                    _add_placemark(f'{layer_name}-LWPOLY', 'Polygon', pts)
                else:
                    _add_placemark(f'{layer_name}-LWPOLY', 'LineString', pts)
                continue

            if etype == 'POLYLINE':
                pts = []
                for v in entity.vertices:
                    pts.append((float(v.dxf.location.x), float(v.dxf.location.y)))
                if len(pts) < 2:
                    continue
                if bool(entity.is_closed) and len(pts) >= 3:
                    if pts[0] != pts[-1]:
                        pts.append(pts[0])
                    _add_placemark(f'{layer_name}-POLYLINE', 'Polygon', pts)
                else:
                    _add_placemark(f'{layer_name}-POLYLINE', 'LineString', pts)
                continue

            stats['unsupported_count'] += 1

        kml_bytes = ET.tostring(kml_root, encoding='utf-8', xml_declaration=True)
        return kml_bytes, stats
    finally:
        if temp_path:
            try:
                os.remove(temp_path)
            except OSError:
                pass


@staff_member_required
def kml_process_convert_view(request):
    """旧 KML 转换页已迁移到 Vue，保留兼容入口。"""
    query_string = request.META.get('QUERY_STRING', '')
    target = '/static/frontend/gis/kml-process-convert'
    if query_string:
        target = f'{target}?{query_string}'
    return redirect(target)


@staff_member_required
@staff_member_required
def ovkml_converter_view(request):
    """旧 OVKML 转换页已迁移到 Vue，保留兼容入口。"""
    query_string = request.META.get('QUERY_STRING', '')
    target = '/static/frontend/gis/ovkml-convert'
    if query_string:
        target = f'{target}?{query_string}'
    return redirect(target)


@staff_member_required
def land_project_management_view(request):
    """旧项目管理页面已迁移到 Vue，保留兼容入口。"""
    return redirect('/static/frontend/projects')


@staff_member_required
def land_project_edit_view(request):
    """旧项目编辑页已迁移到 Vue，保留兼容入口。"""
    project_id = (request.GET.get('project_id') or '').strip()
    if project_id:
        return redirect(f'/static/frontend/projects/{project_id}')
    return redirect('/static/frontend/projects')


@staff_member_required
def land_project_list_api(request):
    """用地项目列表API。"""
    if not is_management_admin(request.user):
        return JsonResponse({'success': False, 'message': '无权限'}, status=403)

    q = (request.GET.get('q') or '').strip()
    status = (request.GET.get('status') or '').strip()
    queryset = LandUseProjectApproval.objects.all()

    if q:
        search_filter = (
            Q(project_name__icontains=q)
            | Q(company_name__icontains=q)
            | Q(final_reply_to_company__icontains=q)
        )
        normalized_date_text = q.replace('/', '-')
        try:
            parsed_date = datetime.strptime(normalized_date_text, '%Y-%m-%d').date()
            search_filter = search_filter | Q(incoming_doc_date=parsed_date)
        except ValueError:
            parsed_date = None
        queryset = queryset.filter(search_filter)
    if status:
        queryset = queryset.filter(status=status)

    rows = []
    for item in queryset.order_by('-receive_date', '-updated_at')[:300]:
        rows.append({
            'id': str(item.id),
            'project_name': item.project_name,
            'company_name': item.company_name,
            'incoming_doc_date': item.incoming_doc_date.isoformat() if item.incoming_doc_date else '',
            'receive_date': item.receive_date.isoformat() if item.receive_date else '',
            'status': item.status,
            'status_label': item.get_status_display(),
            'is_overlap_artifact': item.is_overlap_artifact,
            'updated_at': item.updated_at.strftime('%Y-%m-%d %H:%M') if item.updated_at else '',
        })

    return JsonResponse({'success': True, 'rows': rows})


@staff_member_required
def land_project_detail_api(request, project_id):
    """用地项目详情API。"""
    if not is_management_admin(request.user):
        return JsonResponse({'success': False, 'message': '无权限'}, status=403)

    project = LandUseProjectApproval.objects.filter(id=project_id).first()
    if not project:
        return JsonResponse({'success': False, 'message': '项目不存在'}, status=404)

    photo_rows = [
        {
            'id': photo.id,
            'photo_path': photo.photo_path,
            'photo_url': f"{settings.MEDIA_URL}{photo.photo_path}",
            'uploaded_at': photo.uploaded_at.strftime('%Y-%m-%d %H:%M'),
            'note': photo.note,
        }
        for photo in project.field_photos.all().order_by('-uploaded_at')[:200]
    ]

    operation_logs = [
        {
            'id': item.id,
            'action': item.action,
            'action_label': item.action_label,
            'operator': item.operator.username if item.operator else '系统',
            'payload': item.payload,
            'status_before': item.status_before,
            'status_after': item.status_after,
            'created_at': item.created_at.strftime('%Y-%m-%d %H:%M:%S'),
        }
        for item in project.operation_logs.select_related('operator').all()[:200]
    ]

    return JsonResponse({
        'success': True,
        'data': {
            'id': str(project.id),
            'project_name': project.project_name,
            'company_name': project.company_name,
            'incoming_doc_date': project.incoming_doc_date.isoformat() if project.incoming_doc_date else '',
            'receive_date': project.receive_date.isoformat() if project.receive_date else '',
            'kml_file_path': project.kml_file_path,
            'misc_zip_path': project.misc_zip_path,
            'misc_zip_url': f"/api/land-projects/{project.id}/download-misc-zip/" if project.misc_zip_path else '',
            'is_overlap_artifact': project.is_overlap_artifact,
            'overlapped_relics_info': project.overlapped_relics_info,
            'status': project.status,
            'status_label': project.get_status_display(),
            'field_check_date': project.field_check_date.isoformat() if project.field_check_date else '',
            'shanshan_request_num': project.shanshan_request_num,
            'city_reply_num': project.city_reply_num,
            'archaeology_request_num': project.archaeology_request_num,
            'archaeology_report_path': project.archaeology_report_path,
            'region_approval_num': project.region_approval_num,
            'city_final_reply_num': project.city_final_reply_num,
            'final_reply_to_company': project.final_reply_to_company,
            'controls': get_status_controls(project.status),
            'field_photos': photo_rows,
            'operation_logs': operation_logs,
            'created_at': project.created_at.strftime('%Y-%m-%d %H:%M'),
            'updated_at': project.updated_at.strftime('%Y-%m-%d %H:%M'),
        }
    })


def _load_json_payload(request):
    if not request.body:
        return {}
    try:
        return json.loads(request.body.decode('utf-8'))
    except (json.JSONDecodeError, UnicodeDecodeError):
        raise ValueError('请求体必须是合法JSON')


def _parse_incoming_doc_date(raw_value):
    if not raw_value:
        raise ValueError('incoming_doc_date 必填')
    text = str(raw_value).strip().replace('/', '-')
    try:
        return datetime.strptime(text, '%Y-%m-%d').date()
    except ValueError:
        raise ValueError('incoming_doc_date 格式错误，需为 YYYY-MM-DD')


@csrf_exempt
@require_POST
@staff_member_required
def land_project_create_api(request):
    """收文登记：创建用地项目审批记录。"""
    if not is_management_admin(request.user):
        return JsonResponse({'success': False, 'message': '无权限'}, status=403)

    try:
        payload = _load_json_payload(request)
    except ValueError as exc:
        return JsonResponse({'success': False, 'message': str(exc)}, status=400)

    project_name = (payload.get('project_name') or '').strip()
    company_name = (payload.get('company_name') or '').strip()
    incoming_doc_date_raw = payload.get('incoming_doc_date')
    receive_date = payload.get('receive_date') or timezone.localdate()

    if not project_name or not company_name:
        return JsonResponse({'success': False, 'message': 'project_name/company_name 必填'}, status=400)

    try:
        incoming_doc_date = _parse_incoming_doc_date(incoming_doc_date_raw)
    except ValueError as exc:
        return JsonResponse({'success': False, 'message': str(exc)}, status=400)

    try:
        project = LandUseProjectApproval.objects.create(
            project_name=project_name,
            company_name=company_name,
            incoming_doc_date=incoming_doc_date,
            receive_date=receive_date,
        )
        _record_land_project_operation(
            project=project,
            user=request.user,
            action='create',
            payload={
                'project_name': project_name,
                'company_name': company_name,
                'incoming_doc_date': incoming_doc_date.isoformat(),
            },
            status_before='',
            status_after=project.status,
        )
    except (OperationalError, ProgrammingError) as exc:
        logger.exception('创建项目失败，疑似数据库结构未就绪: %s', exc)
        return JsonResponse(
            {
                'success': False,
                'message': '数据库结构未就绪，请在服务器执行 python manage.py migrate 后重试',
                'error_type': exc.__class__.__name__,
            },
            status=500,
        )
    except Exception as exc:
        logger.exception('创建项目失败: %s', exc)
        return JsonResponse(
            {
                'success': False,
                'message': '创建项目失败，请查看服务端日志',
                'error_type': exc.__class__.__name__,
            },
            status=500,
        )

    return JsonResponse({'success': True, 'project_id': str(project.id), 'status': project.status})


@csrf_exempt
@require_POST
@staff_member_required
def land_project_upload_api(request, project_id):
    """
    文件自动归档：
    - kml: projects/{year}/{项目名}/kml/
    - misc_zip: projects/{year}/{项目名}/misc/
    - field_photo: projects/{year}/{项目名}/field_checks/
    - archaeology_report: projects/{year}/{项目名}/archaeology/
    """
    if not is_management_admin(request.user):
        return JsonResponse({'success': False, 'message': '无权限'}, status=403)

    project = LandUseProjectApproval.objects.filter(id=project_id).first()
    if not project:
        return JsonResponse({'success': False, 'message': '项目不存在'}, status=404)

    upload_file = request.FILES.get('file')
    file_type = (request.POST.get('file_type') or '').strip()
    if not upload_file:
        return JsonResponse({'success': False, 'message': '未接收到上传文件'}, status=400)

    filename = os.path.basename(upload_file.name or 'upload.bin')

    if file_type == 'kml':
        if not filename.lower().endswith(('.kml', '.kmz', '.ovkml', '.ovkmz')):
            return JsonResponse({'success': False, 'message': 'KML文件类型不正确'}, status=400)
        status_before = project.status
        relative_path = build_project_media_path(project, 'kml', filename)
        saved_path = default_storage.save(relative_path, upload_file)
        project.kml_file_path = saved_path
        project.save(update_fields=['kml_file_path', 'updated_at'])
        _record_land_project_operation(
            project=project,
            user=request.user,
            action='upload_kml',
            payload={
                'file_type': file_type,
                'original_filename': filename,
                'saved_path': saved_path,
            },
            status_before=status_before,
            status_after=project.status,
        )
        return JsonResponse({'success': True, 'file_path': saved_path})

    if file_type == 'field_photo':
        status_before = project.status
        relative_path = build_project_media_path(project, 'field_checks', filename)
        saved_path = default_storage.save(relative_path, upload_file)
        photo = LandUseProjectFieldPhoto.objects.create(
            project=project,
            photo_path=saved_path,
            note=(request.POST.get('note') or '').strip(),
        )
        _record_land_project_operation(
            project=project,
            user=request.user,
            action='upload_field_photo',
            payload={
                'file_type': file_type,
                'original_filename': filename,
                'saved_path': saved_path,
                'photo_id': photo.id,
            },
            status_before=status_before,
            status_after=project.status,
        )
        return JsonResponse({'success': True, 'photo_id': photo.id, 'file_path': saved_path})

    if file_type == 'misc_zip':
        if not filename.lower().endswith('.zip'):
            return JsonResponse({'success': False, 'message': '杂项文件仅支持ZIP格式'}, status=400)
        status_before = project.status
        relative_path = build_project_media_path(project, 'misc', filename)
        saved_path = default_storage.save(relative_path, upload_file)
        project.misc_zip_path = saved_path
        project.save(update_fields=['misc_zip_path', 'updated_at'])
        _record_land_project_operation(
            project=project,
            user=request.user,
            action='upload_misc_zip',
            payload={
                'file_type': file_type,
                'original_filename': filename,
                'saved_path': saved_path,
            },
            status_before=status_before,
            status_after=project.status,
        )
        return JsonResponse({'success': True, 'file_path': saved_path})

    if file_type == 'archaeology_report':
        if not filename.lower().endswith('.pdf'):
            return JsonResponse({'success': False, 'message': '考古调查报告仅支持PDF'}, status=400)
        status_before = project.status
        relative_path = build_project_media_path(project, 'archaeology', filename)
        saved_path = default_storage.save(relative_path, upload_file)
        project.archaeology_report_path = saved_path
        project.save(update_fields=['archaeology_report_path', 'updated_at'])
        _record_land_project_operation(
            project=project,
            user=request.user,
            action='upload_archaeology_report',
            payload={
                'file_type': file_type,
                'original_filename': filename,
                'saved_path': saved_path,
            },
            status_before=status_before,
            status_after=project.status,
        )
        return JsonResponse({'success': True, 'file_path': saved_path})

    return JsonResponse({'success': False, 'message': 'file_type 必须为 kml/misc_zip/field_photo/archaeology_report'}, status=400)


@staff_member_required
def land_project_download_misc_zip_api(request, project_id):
    """下载项目杂项ZIP文件。"""
    if not is_management_admin(request.user):
        return JsonResponse({'success': False, 'message': '无权限'}, status=403)

    project = LandUseProjectApproval.objects.filter(id=project_id).first()
    if not project:
        return JsonResponse({'success': False, 'message': '项目不存在'}, status=404)
    if not project.misc_zip_path:
        return JsonResponse({'success': False, 'message': '当前项目未上传杂项ZIP'}, status=404)
    if not default_storage.exists(project.misc_zip_path):
        return JsonResponse({'success': False, 'message': '杂项ZIP文件不存在或已被移除'}, status=404)

    try:
        file_handler = default_storage.open(project.misc_zip_path, 'rb')
    except Exception:
        logger.exception('打开杂项ZIP失败: project_id=%s, path=%s', project_id, project.misc_zip_path)
        return JsonResponse({'success': False, 'message': '文件读取失败'}, status=500)

    download_name = os.path.basename(project.misc_zip_path) or f'{project.project_name}_misc.zip'
    return FileResponse(file_handler, as_attachment=True, filename=download_name, content_type='application/zip')


@csrf_exempt
@require_POST
@staff_member_required
def verify_project_spatial_safety_api(request, project_id):
    """触发KML与文物保护范围/建控地带叠加核验。"""
    if not is_management_admin(request.user):
        return JsonResponse({'success': False, 'message': '无权限'}, status=403)

    project = LandUseProjectApproval.objects.filter(id=project_id).first()
    if not project:
        return JsonResponse({'success': False, 'message': '项目不存在'}, status=404)

    status_before = project.status

    try:
        result = verify_project_spatial_safety(project_id)
        _record_land_project_operation(
            project=project,
            user=request.user,
            action='verify_spatial_safety',
            payload={
                'is_overlap_artifact': result.get('is_overlap_artifact', False),
                'overlapped_count': len(result.get('overlapped_relics_info') or []),
            },
            status_before=status_before,
            status_after=result.get('status', project.status),
        )
    except ValueError as exc:
        return JsonResponse({'success': False, 'message': str(exc)}, status=400)
    except Exception:
        logger.exception('项目空间核验失败: project_id=%s', project_id)
        return JsonResponse({'success': False, 'message': '空间核验失败，请检查KML与空间数据'}, status=500)

    return JsonResponse({'success': True, 'data': result})


@staff_member_required
def land_project_next_doc_num_api(request):
    """文号推荐：根据年度历史数据推荐下一个鄯文旅字文号。"""
    if not is_management_admin(request.user):
        return JsonResponse({'success': False, 'message': '无权限'}, status=403)

    raw_year = (request.GET.get('year') or '').strip()
    year = int(raw_year) if raw_year.isdigit() else timezone.localdate().year
    value = LandUseProjectApproval.suggest_next_shanshan_num(year)
    return JsonResponse({'success': True, 'year': year, 'next_doc_num': value})


@csrf_exempt
@require_POST
@staff_member_required
def land_project_workflow_action_api(request, project_id):
    """状态机强控流转入口。"""
    if not is_management_admin(request.user):
        return JsonResponse({'success': False, 'message': '无权限'}, status=403)

    project = LandUseProjectApproval.objects.filter(id=project_id).first()
    if not project:
        return JsonResponse({'success': False, 'message': '项目不存在'}, status=404)

    try:
        payload = _load_json_payload(request)
        action = (payload.get('action') or '').strip()
        status_before = project.status
        result = apply_workflow_action(project, action, payload)
        _record_land_project_operation(
            project=project,
            user=request.user,
            action=action,
            payload=_build_payload_doc_nums(payload),
            status_before=status_before,
            status_after=result.get('status', project.status),
        )
    except ValueError as exc:
        return JsonResponse({'success': False, 'message': str(exc)}, status=400)
    except Exception:
        logger.exception('项目流程动作执行失败: project_id=%s', project_id)
        return JsonResponse({'success': False, 'message': '流程动作执行失败'}, status=500)

    return JsonResponse({'success': True, 'data': result})


@staff_member_required
def land_project_controls_api(request, project_id):
    """前端按钮动态启禁：返回当前状态对应的操作许可。"""
    if not is_management_admin(request.user):
        return JsonResponse({'success': False, 'message': '无权限'}, status=403)

    project = LandUseProjectApproval.objects.filter(id=project_id).first()
    if not project:
        return JsonResponse({'success': False, 'message': '项目不存在'}, status=404)

    return JsonResponse({
        'success': True,
        'project_id': str(project.id),
        'status': project.status,
        'status_label': project.get_status_display(),
        'controls': get_status_controls(project.status),
    })


@staff_member_required
def heritage_dashboard_view(request):
    """旧统计页已迁移到 Vue，保留兼容入口。"""
    return redirect('/static/frontend/heritage/stats')


def _extract_township_name(address):
    if not address:
        return ''

    text = str(address).strip()
    for keyword, standard_name in TOWNSHIP_NORMALIZATION_RULES:
        if keyword in text:
            return standard_name

    match = re.search(r'鄯善县(?:吐鲁番市鄯善县)*(?:东北)?([\u4e00-\u9fa5]{1,12}?(?:回族乡|乡|镇|街道))', text)
    if match:
        return match.group(1)

    return ''


def _to_township_full_name(township_name):
    if not township_name:
        return ''
    return township_name


def _apply_kanerjing_filter(queryset, kanerjing_scope):
    if kanerjing_scope == 'only':
        return HeritageSite.filter_kanerjing(queryset)
    if kanerjing_scope == 'exclude':
        return HeritageSite.exclude_kanerjing(queryset)
    return queryset


def _resolve_heritage_stats_source(source):
    source_key = (source or 'auto').strip().lower()
    if source_key == 'legacy':
        return {
            'source': 'legacy',
            'model': HeritageSite,
            'category_field': 'category',
            'level_field': 'level',
            'township_field': '',
            'address_field': 'address',
            'supports_kanerjing': True,
        }

    if source_key == 'immovable':
        return {
            'source': 'immovable',
            'model': ImmovableHeritage,
            'category_field': 'category',
            'level_field': 'protection_level',
            'township_field': 'township',
            'address_field': 'address',
            'supports_kanerjing': False,
        }

    # 统一口径：auto 模式固定回退到 HeritageSite（legacy）。
    if source_key == 'auto':
        return {
            'source': 'legacy',
            'model': HeritageSite,
            'category_field': 'category',
            'level_field': 'level',
            'township_field': '',
            'address_field': 'address',
            'supports_kanerjing': True,
        }

    return {
        'source': 'legacy',
        'model': HeritageSite,
        'category_field': 'category',
        'level_field': 'level',
        'township_field': '',
        'address_field': 'address',
        'supports_kanerjing': True,
    }


def _get_field_choices_map(model_cls, field_name):
    if not field_name:
        return {}
    try:
        return dict(model_cls._meta.get_field(field_name).choices or [])
    except Exception:
        return {}


def _build_group_rows(queryset, group_field, choices_map):
    raw_rows = queryset.values(group_field).annotate(count=Count('id')).order_by('-count')
    rows = []
    for item in raw_rows:
        value = item.get(group_field)
        normalized = str(value).strip() if value is not None else ''
        label = choices_map.get(value) or choices_map.get(normalized) or normalized or '未标注'
        rows.append({'value': normalized, 'label': label, 'count': item['count']})
    return rows


@staff_member_required
def heritage_classification_stats_api(request):
    """文物分类统计 API：按数据库真实字段自动分组统计，并保持旧结构兼容。"""
    category = request.GET.get('category', '').strip()
    level = request.GET.get('level', '').strip()
    township = request.GET.get('township', '').strip()
    address_keyword = request.GET.get('address_keyword', '').strip()
    kanerjing_scope = request.GET.get('kanerjing_scope', 'all').strip()
    group_by = request.GET.get('group_by', 'category').strip()
    source = request.GET.get('source', 'auto').strip()

    source_config = _resolve_heritage_stats_source(source)
    model_cls = source_config['model']
    category_field = source_config['category_field']
    level_field = source_config['level_field']
    township_field = source_config['township_field']
    address_field = source_config['address_field']

    queryset = model_cls.objects.all()
    if category:
        queryset = queryset.filter(**{category_field: category})
    if level:
        queryset = queryset.filter(**{level_field: level})
    if township:
        if township_field:
            queryset = queryset.filter(**{f'{township_field}__icontains': township})
        else:
            township_keywords = TOWNSHIP_STANDARD_TO_KEYWORDS.get(township, {township})
            township_query = Q()
            for keyword in township_keywords:
                township_query |= Q(address__icontains=keyword)
            queryset = queryset.filter(township_query)
    if address_keyword:
        queryset = queryset.filter(**{f'{address_field}__icontains': address_keyword})

    if source_config['supports_kanerjing']:
        queryset = _apply_kanerjing_filter(queryset, kanerjing_scope)
    else:
        kanerjing_scope = 'all'

    group_key = (group_by or 'category').strip().lower()
    if group_key in {'level', 'protection_level'}:
        group_field = level_field
        group_key = 'level'
    elif group_key in {'category'}:
        group_field = category_field
        group_key = 'category'
    elif group_key in {'heritage_type'} and hasattr(model_cls, 'HERITAGE_TYPE_CHOICES'):
        group_field = 'heritage_type'
        group_key = 'heritage_type'
    elif group_key == 'township':
        group_field = 'township' if township_field else ''
        group_key = 'township'
    else:
        group_field = category_field
        group_key = 'category'

    if group_key == 'township' and not group_field:
        township_counter = {}
        for item in queryset.values(address_field):
            township_name = _extract_township_name(item.get(address_field))
            key = township_name or '未标注乡镇'
            township_counter[key] = township_counter.get(key, 0) + 1
        sorted_items = sorted(township_counter.items(), key=lambda x: x[1], reverse=True)
        rows = [
            {
                'value': item[0],
                'label': '未标注乡镇' if item[0] == '未标注乡镇' else _to_township_full_name(item[0]),
                'count': item[1],
            }
            for item in sorted_items
        ]
    else:
        choices_map = _get_field_choices_map(model_cls, group_field)
        rows = _build_group_rows(queryset, group_field, choices_map)

    labels = [item['label'] for item in rows]
    data = [item['count'] for item in rows]

    return JsonResponse({
        'success': True,
        'labels': labels,
        'data': data,
        'rows': rows,
        'total': queryset.count(),
        'group_by': group_key,
        'group_by_field': group_field or 'township',
        'source': source_config['source'],
        'storage_fields': {
            'category': category_field,
            'level': level_field,
            'township': township_field or 'address(extracted)',
        },
        'kanerjing_scope': kanerjing_scope if kanerjing_scope in {'all', 'only', 'exclude'} else 'all',
    })


@staff_member_required
def heritage_stats_api(request):
    """统计数据 API 端点"""
    total = HeritageSite.objects.count()
    national = HeritageSite.objects.filter(level='GB').count()
    regional = HeritageSite.objects.filter(level='SB').count()
    county = HeritageSite.objects.filter(level='XB').count()
    
    return JsonResponse({
        'total': total,
        'national': national,
        'regional': regional,
        'county': county,
    })


@csrf_exempt
@require_POST
def dem_elevation_lookup_api(request):
    """DEM 海拔反查 API：接收经纬度，返回 SRTM 30m 高程。"""
    try:
        payload = json.loads(request.body.decode('utf-8')) if request.body else {}
    except (json.JSONDecodeError, UnicodeDecodeError):
        return JsonResponse(
            {
                'success': False,
                'message': '请求体必须是合法 JSON',
                'elevation': None,
            },
            status=400,
        )

    longitude = payload.get('longitude')
    latitude = payload.get('latitude')

    try:
        lon = float(longitude)
        lat = float(latitude)
    except (TypeError, ValueError):
        return JsonResponse(
            {
                'success': False,
                'message': 'longitude/latitude 必须是数字',
                'elevation': None,
            },
            status=400,
        )

    if not (-180 <= lon <= 180 and -90 <= lat <= 90):
        return JsonResponse(
            {
                'success': False,
                'message': '经纬度超出有效范围',
                'elevation': None,
            },
            status=400,
        )

    tile_name = describe_tile(lon, lat)

    prefer_cached_only = bool(payload.get('prefer_cached_only', True))

    try:
        elevation = get_dem_elevation(
            lon,
            lat,
            prefer_cached_only=prefer_cached_only,
            trigger_background_download=True,
        )
    except Exception:
        # 双保险兜底：任何 DEM 层异常都不影响主业务流程。
        import logging

        logging.getLogger(__name__).exception(
            'DEM 反查发生未捕获异常，lon=%s, lat=%s, tile=%s',
            lon,
            lat,
            tile_name,
        )
        elevation = None

    return JsonResponse(
        {
            'success': True,
            'longitude': lon,
            'latitude': lat,
            'tile': tile_name,
            'elevation': elevation,
            'prefer_cached_only': prefer_cached_only,
            'cache_hit': elevation is not None,
            'source': 'SRTMGL1(OpenTopography)',
            'fallback_hint': 'elevation 为 null 时，请前端降级使用设备海拔',
        }
    )

@staff_member_required
def heritage_stats_by_category_api(request):
    """按类别统计的 API 端点"""
    category_stats = HeritageSite.objects.values('category').annotate(count=Count('id'))
    
    labels = []
    data = []
    
    for stat in category_stats:
        category_code = stat['category']
        # 获取中文标签
        category_name = dict(HeritageSite.CATEGORY_CHOICES).get(category_code, category_code)
        labels.append(category_name)
        data.append(stat['count'])
    
    return JsonResponse({
        'labels': labels,
        'data': data,
    })

@staff_member_required
def kanerjing_list_view(request):
    """旧坎儿井专项管理页已迁移到 Vue，保留兼容入口。"""
    return redirect('/static/frontend/heritage/kanerjing')

@staff_member_required
def kanerjing_stats_api(request):
    """坎儿井统计 API - 基于名称包含'坎儿井'进行筛选"""
    kanerjing_sites = HeritageSite.filter_kanerjing()
    total = kanerjing_sites.count()
    
    # 按等级统计
    level_breakdown = {}
    for level_code, level_name in HeritageSite.LEVEL_CHOICES:
        count = kanerjing_sites.filter(level=level_code).count()
        level_breakdown[level_name] = count
    
    # 地址分组按镇/乡统计，只保留“xx镇/xx乡”层级。
    def normalize_address(address):
        if not address:
            return '未标注镇乡'

        text = str(address).strip()
        if not text:
            return '未标注镇乡'

        # 复用统一乡镇提取规则，避免把完整行政区划误当作乡镇标签。
        township_name = _extract_township_name(text)
        if township_name:
            return township_name

        # 兜底：取地址中最后一个“xx镇/xx乡/xx回族乡/xx街道”。
        fallback_matches = re.findall(r'([\u4e00-\u9fa5]{1,12}(?:回族乡|乡|镇|街道))', text)
        if fallback_matches:
            return fallback_matches[-1]

        return '未标注镇乡'

    address_counter = {}
    for address in kanerjing_sites.values_list('address', flat=True):
        normalized = normalize_address(address)
        address_counter[normalized] = address_counter.get(normalized, 0) + 1

    address_distribution = [
        {'township': addr, 'address': addr, 'count': count}
        for addr, count in sorted(address_counter.items(), key=lambda item: (-item[1], item[0]))[:10]
    ]
    
    return JsonResponse({
        'total': total,
        'by_level': level_breakdown,
        'top_addresses': address_distribution,
    })

@staff_member_required
def kanerjing_import_check_view(request):
    """坎儿井导入检查页面已下线，统一回到坎儿井专项管理。"""
    return redirect('/static/frontend/heritage/kanerjing')


# 自定义密码修改完成视图
from django.contrib.auth.views import PasswordChangeDoneView
from .signals import mark_password_changed


class CustomPasswordChangeDoneView(PasswordChangeDoneView):
    """自定义密码修改完成视图 - 标记用户已修改密码"""
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        # 标记用户已修改密码
        mark_password_changed(self.request.user)
        
        # 添加自定义信息
        from django.contrib import messages
        messages.success(
            self.request,
            '✓ 密码修改成功！您的账户更安全了。'
        )
        
        return context


@staff_member_required
def inspection_mobile_add_view(request):
    """旧手机巡查新增页已迁移到 Vue，保留兼容入口。"""
    return redirect('/static/frontend/heritage/inspections/new')


@staff_member_required
def inspection_mobile_list_view(request):
    """旧手机巡查列表已迁移到 Vue，保留兼容入口。"""
    return redirect('/static/frontend/heritage/inspections')


def app_showcase_view(request):
    """鄯善文保 App 展示与下载页（公开访问）"""
    context = {
        'download_url': 'https://share.fnnas.net/s/49a5a7b485784cf4a1',
        'qr_code_url': 'https://api.qrserver.com/v1/create-qr-code/?size=280x280&data=https%3A%2F%2Fshare.fnnas.net%2Fs%2F49a5a7b485784cf4a1',
        'features': [
            '文物点巡查上报：支持现场拍照、位置记录与问题描述，提升巡查效率。',
            '巡查记录管理：随时查看历史巡查内容，支持按时间快速追溯。',
            '关联中心协同：重点任务与相关文物信息关联展示，便于统一处置。',
            '个人账户中心：看护员可管理个人信息与使用入口，操作清晰。',
        ],
        'screenshots': [
            {'name': '登录页面', 'file': 'app_showcase/登录页面.jpg'},
            {'name': '工作台', 'file': 'app_showcase/工作台.jpg'},
            {'name': '巡查上报', 'file': 'app_showcase/巡查上报.jpg'},
            {'name': '巡查记录', 'file': 'app_showcase/巡查记录.jpg'},
            {'name': '关联中心', 'file': 'app_showcase/关联中心.jpg'},
            {'name': '我的账户', 'file': 'app_showcase/我的账户.jpg'},
        ],
    }
    return render(request, 'public/app_showcase.html', context)


# ─────────────────────────────────────────────────────────────────────────────
# 不可移动文物采集入口（旧平台兼容）
# ─────────────────────────────────────────────────────────────────────────────
from django.contrib.auth.decorators import login_required


def _format_coord_point_lines(points):
    if not points:
        return ['无']

    type_map = {
        'boundary': '边界点',
        'marker': '标志点',
        'other': '其他',
    }
    lines = []
    for idx, point in enumerate(points, start=1):
        p_type = type_map.get(str(point.get('type') or ''), '其他')
        lon = point.get('longitude', '')
        lat = point.get('latitude', '')
        alt = point.get('altitude')
        desc = str(point.get('description') or '').strip() or '无'
        src = str(point.get('sourceTag') or '').strip()
        remark = str(point.get('remark') or '').strip()

        alt_text = f"{alt:.2f}" if isinstance(alt, (int, float)) else '无'
        line = f"{idx}. [{p_type}] 经度 {lon}，纬度 {lat}，海拔 {alt_text}m，说明：{desc}"
        if src:
            line = f"{line}，来源：{src}"
        if remark:
            line = f"{line}，备注：{remark}"
        lines.append(line)

    return lines


def _build_coord_points_display(points):
    type_map = {
        'boundary': '边界点',
        'marker': '标志点',
        'other': '其他',
    }
    rows = []
    if not isinstance(points, list):
        return rows

    for idx, point in enumerate(points, start=1):
        if not isinstance(point, dict):
            continue
        try:
            lon = float(point.get('longitude'))
            lat = float(point.get('latitude'))
        except (TypeError, ValueError):
            continue

        alt = point.get('altitude')
        alt_text = '——'
        if alt not in (None, ''):
            try:
                alt_text = f"{float(alt):.2f}"
            except (TypeError, ValueError):
                alt_text = '——'

        rows.append({
            'index': idx,
            'type_label': type_map.get(str(point.get('type') or ''), '其他'),
            'longitude': f"{lon:.8f}",
            'latitude': f"{lat:.8f}",
            'altitude': alt_text,
            'description': str(point.get('description') or '').strip() or '——',
            'source_tag': str(point.get('sourceTag') or '').strip(),
            'remark': str(point.get('remark') or '').strip() or '——',
        })
    return rows


@login_required
def heritage_collect_view(request):
    """旧采集页面已下线，统一跳转至 Vue 文物采集页。"""
    return redirect('/static/frontend/collect/immovable')


# ─────────────────────────────────────────────────────────────────────────────
# 采集登记表预览视图（仅登录用户可访问）
# ─────────────────────────────────────────────────────────────────────────────
import math as _math


def _decimal_to_dms(decimal_deg, is_longitude: bool) -> str:
    """
    将十进制度转换为"度°分′秒″"格式（中文方向符号）。
    例：90.335167 → 东经 90°20′06.60″
    """
    try:
        val = float(decimal_deg)
    except (TypeError, ValueError):
        return "——"

    direction = ""
    if is_longitude:
        direction = "东经" if val >= 0 else "西经"
    else:
        direction = "北纬" if val >= 0 else "南纬"

    val = abs(val)
    degrees = int(val)
    minutes_float = (val - degrees) * 60
    minutes = int(minutes_float)
    seconds = (minutes_float - minutes) * 60

    return f"{direction} {degrees}°{minutes:02d}′{seconds:05.2f}″"


def _safe_file_url(file_field) -> str:
    """安全获取文件 URL，避免坏数据或空文件导致预览页 500。"""
    try:
        if file_field and getattr(file_field, "name", ""):
            return file_field.url
    except Exception:
        return ""
    return ""


def _display_user_name(user) -> str:
    """统一处理用户显示名称，避免模板对空对象链式取值。"""
    if not user:
        return "——"
    full_name = user.get_full_name() if hasattr(user, "get_full_name") else ""
    return full_name or getattr(user, "username", "——") or "——"


def _chunk_items(items, chunk_size):
    """按固定大小切分列表，便于登记表照片分页。"""
    if chunk_size <= 0:
        return [items]
    return [items[index:index + chunk_size] for index in range(0, len(items), chunk_size)]


@login_required
def heritage_detail_preview_view(request, pk):
    """
    不可移动文物采集登记表预览页面。

    - 任意已登录用户均可访问（字段只读，无修改功能）。
    - 提供打印友好的 A4 布局，可直接从浏览器打印为 PDF。
    - URL：/mobile/collect/<int:pk>/preview/
    """
    from .models import ImmovableHeritage, HeritagePhoto
    from django.shortcuts import get_object_or_404

    heritage = (
        ImmovableHeritage.objects.select_related("collector", "input_by", "reviewer")
        .prefetch_related("photos")
        .filter(pk=pk)
        .first()
    )

    if heritage:
        # 照片：封面 + 其余（最多展示 8 张附图，避免撑破页面）
        all_photos = list(heritage.photos.exclude(image="").order_by("-is_cover", "-shot_at", "-uploaded_at"))
        cover_photo = next((p for p in all_photos if p.is_cover), None) or (all_photos[0] if all_photos else None)
        other_photos = [p for p in all_photos if p != cover_photo][:8]
    else:
        # 兼容旧档案库：当 pk 来自 HeritageSite 时，回退到基础档案并构造预览对象。
        site = get_object_or_404(HeritageSite, pk=pk)
        heritage = SimpleNamespace(
            survey_code=site.sip_code or f"HS-{site.id}",
            previous_survey_code="",
            collected_at=getattr(site, "created_at", None),
            reviewed_at=None,
            name=site.name,
            former_name="",
            era="——",
            category=site.category,
            heritage_type="",
            province="新疆维吾尔自治区",
            city="吐鲁番市",
            county="鄯善县",
            township="",
            village="",
            address=site.address or "",
            coordinate_system="CGCS2000",
            longitude=site.longitude,
            latitude=site.latitude,
            altitude=None,
            area=None,
            preservation_status="一般",
            is_disappeared=False,
            disappear_reason="",
            is_relocated=False,
            relocation_note="",
            damage_cause="",
            threat_factors="",
            ownership="state",
            ownership_detail="",
            user_unit="",
            management_unit="",
            manager=site.manager or "",
            protection_level=site.level,
            protection_announced_batch="",
            protection_announced_date=None,
            has_marker_stele=False,
            has_protection_zone_announced=False,
            has_construction_control_zone_announced=False,
            description=site.description or "",
            remarks="",
            collector=None,
            input_by=None,
            reviewer=None,
            coord_list=[],
            get_category_display=site.get_category_display,
            get_coordinate_system_display=lambda: "2000国家大地坐标系（CGCS2000）",
            get_protection_level_display=site.get_level_display,
            get_ownership_display=lambda: "国有",
        )
        all_photos = []
        cover_photo = None
        other_photos = []

    # 度分秒在视图层计算，保持模板简洁
    lon_dms = _decimal_to_dms(heritage.longitude, is_longitude=True)
    lat_dms = _decimal_to_dms(heritage.latitude,  is_longitude=False)
    coord_points = _build_coord_points_display(heritage.coord_list)
    cover_photo_url = _safe_file_url(cover_photo.image if cover_photo else None)
    other_photo_items = [
        {
            "url": _safe_file_url(photo.image),
            "caption": photo.caption,
            "photo_type_display": photo.get_photo_type_display(),
            "shot_at": photo.shot_at,
        }
        for photo in other_photos
        if _safe_file_url(photo.image)
    ]
    other_photo_pages = _chunk_items(other_photo_items, 4)
    collect_unit = "鄯善县文化体育广播电视和旅游局（文物局）"

    view_mode = (request.GET.get("mode") or "view").strip().lower()
    if view_mode not in {"view", "print"}:
        view_mode = "view"

    context = {
        "heritage":             heritage,
        "cover_photo":          cover_photo,
        "cover_photo_url":      cover_photo_url,
        "other_photos":         other_photos,
        "other_photo_items":    other_photo_items,
        "other_photo_pages":    other_photo_pages,
        "photos":               all_photos,
        "lon_dms":              lon_dms,
        "lat_dms":              lat_dms,
        "coord_points":         coord_points,
        "collector_display":    _display_user_name(heritage.collector),
        "input_by_display":     _display_user_name(heritage.input_by),
        "reviewer_display":     _display_user_name(heritage.reviewer),
        "collect_unit":         collect_unit,
        "preservation_choices": ImmovableHeritage.PRESERVATION_STATUS_CHOICES,
        "protection_choices":   ImmovableHeritage.PROTECTION_LEVEL_CHOICES,
        "ownership_choices":    ImmovableHeritage.OWNERSHIP_CHOICES,
        "view_mode":            view_mode,
    }
    return render(request, "public/detail_preview.html", context)


def _set_cell_text(cell, text, *, bold=False, align=WD_PARAGRAPH_ALIGNMENT.LEFT, font_size=14):
    """统一设置表格单元格文本样式。"""
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run("" if text is None else str(text))
    run.bold = bold
    run.font.size = Pt(font_size)
    run.font.name = "宋体"


def _set_table_column_widths(table, widths_mm):
    """设置表格列宽，避免 Word 自动拉伸导致打印错位。"""
    for row in table.rows:
        for idx, width in enumerate(widths_mm):
            row.cells[idx].width = Mm(width)


def _scaled_size_mm(img_bytes, max_w_mm=70.0, max_h_mm=52.0):
    """按比例缩放图片，返回毫米宽高。"""
    with PILImage.open(io.BytesIO(img_bytes)) as im:
        px_w, px_h = im.size
    if px_w <= 0 or px_h <= 0:
        return max_w_mm, max_h_mm

    ratio = px_w / px_h
    width_mm = max_w_mm
    height_mm = width_mm / ratio
    if height_mm > max_h_mm:
        height_mm = max_h_mm
        width_mm = height_mm * ratio
    return width_mm, height_mm


def _insert_photos_into_cell(cell, photos):
    """
    将现场照片插入单元格。
    优先封面图，其余最多补充 2 张，避免撑破表格。
    """
    cell.text = ""
    if not photos:
        _set_cell_text(cell, "暂无现场照片", align=WD_PARAGRAPH_ALIGNMENT.CENTER, font_size=14)
        return

    ordered = sorted(photos, key=lambda p: (not p.is_cover, p.uploaded_at or timezone.now()))
    selected = ordered[:3]

    for index, photo in enumerate(selected):
        try:
            with photo.image.open("rb") as image_file:
                img_bytes = image_file.read()
        except Exception:
            p_err = cell.add_paragraph()
            p_err.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            p_err.add_run(f"照片{index + 1}读取失败")
            continue

        width_mm, height_mm = _scaled_size_mm(img_bytes, max_w_mm=68.0, max_h_mm=50.0)
        stream = io.BytesIO(img_bytes)

        p_img = cell.add_paragraph()
        p_img.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p_img.add_run().add_picture(stream, width=Mm(width_mm), height=Mm(height_mm))

        p_caption = cell.add_paragraph()
        p_caption.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        caption = f"图{index + 1}  {photo.get_photo_type_display()}"
        if photo.caption:
            caption = f"{caption}：{photo.caption}"
        run_caption = p_caption.add_run(caption)
        run_caption.font.size = Pt(14)
        run_caption.font.name = "宋体"


def _safe_docx_fragment(text):
    raw = (text or '').strip()
    if not raw:
        return 'unknown'
    cleaned = re.sub(r'[\\/:*?"<>|]+', '_', raw)
    cleaned = re.sub(r'\s+', '_', cleaned)
    return cleaned[:80]


def _build_immovable_heritage_docx_stream(heritage):
    """生成单条不可移动文物采集登记表 DOCX，返回 (BytesIO, filename)。"""
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(25.4)
    section.right_margin = Mm(25.4)
    section.top_margin = Mm(25.4)
    section.bottom_margin = Mm(25.4)

    p_title = document.add_paragraph()
    p_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run_title = p_title.add_run("鄯善县不可移动文物采集登记表")
    run_title.bold = True
    run_title.font.size = Pt(22)
    run_title.font.name = "黑体"

    p_code = document.add_paragraph()
    p_code.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run_code = p_code.add_run(f"采集编号：{heritage.survey_code}")
    run_code.bold = True
    run_code.font.size = Pt(14)
    run_code.font.name = "宋体"

    table = document.add_table(rows=17, cols=8)
    table.style = "Table Grid"
    _set_table_column_widths(table, [18, 24, 14, 24, 14, 24, 14, 27])

    _set_cell_text(table.cell(0, 0).merge(table.cell(0, 7)), "一、基本信息", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(1, 0), "名称", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(1, 1).merge(table.cell(1, 3)), heritage.name)
    _set_cell_text(table.cell(1, 4), "曾用名/别名", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(1, 5).merge(table.cell(1, 7)), heritage.former_name or "——")

    _set_cell_text(table.cell(2, 0), "时代", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(2, 1), heritage.era)
    _set_cell_text(table.cell(2, 2), "类别", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(2, 3), heritage.get_category_display())
    _set_cell_text(table.cell(2, 4), "类型", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(2, 5).merge(table.cell(2, 7)), heritage.heritage_type or "——")

    _set_cell_text(table.cell(3, 0).merge(table.cell(3, 7)), "二、地理位置", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(4, 0), "省/自治区", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(4, 1), heritage.province or "——")
    _set_cell_text(table.cell(4, 2), "市/州", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(4, 3), heritage.city or "——")
    _set_cell_text(table.cell(4, 4), "县/区", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(4, 5).merge(table.cell(4, 7)), heritage.county or "——")

    _set_cell_text(table.cell(5, 0), "乡镇/街道", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(5, 1).merge(table.cell(5, 3)), heritage.township or "——")
    _set_cell_text(table.cell(5, 4), "村/社区", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(5, 5).merge(table.cell(5, 7)), heritage.village or "——")

    _set_cell_text(table.cell(6, 0), "详细地址", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(6, 1).merge(table.cell(6, 7)), heritage.address or "——")

    _set_cell_text(table.cell(7, 0), "坐标系", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(7, 1), heritage.get_coordinate_system_display())
    _set_cell_text(table.cell(7, 2), "经度", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(7, 3), f"{heritage.longitude:.8f}", font_size=14)
    _set_cell_text(table.cell(7, 4), "纬度", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(7, 5), f"{heritage.latitude:.8f}", font_size=14)
    _set_cell_text(table.cell(7, 6), "海拔(m)", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(7, 7), f"{heritage.altitude:.2f}" if heritage.altitude is not None else "——", font_size=14)

    _set_cell_text(table.cell(8, 0), "占地面积", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(8, 1).merge(table.cell(8, 3)), f"{heritage.area:.2f} 平方米" if heritage.area else "——")
    _set_cell_text(table.cell(8, 4), "经度(度分秒)", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER, font_size=14)
    _set_cell_text(table.cell(8, 5), _decimal_to_dms(heritage.longitude, True), font_size=14)
    _set_cell_text(table.cell(8, 6), "纬度(度分秒)", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER, font_size=14)
    _set_cell_text(table.cell(8, 7), _decimal_to_dms(heritage.latitude, False), font_size=14)

    _set_cell_text(table.cell(9, 0).merge(table.cell(9, 7)), "三、现状与保护", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(10, 0), "保存现状", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(10, 1), heritage.preservation_status)
    _set_cell_text(table.cell(10, 2), "权属", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(10, 3), heritage.get_ownership_display())
    _set_cell_text(table.cell(10, 4), "保护级别", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(10, 5).merge(table.cell(10, 7)), heritage.get_protection_level_display())

    _set_cell_text(table.cell(11, 0), "破坏原因", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(11, 1).merge(table.cell(11, 3)), heritage.damage_cause or "无")
    _set_cell_text(table.cell(11, 4), "威胁因素", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(11, 5).merge(table.cell(11, 7)), heritage.threat_factors or "无")

    _set_cell_text(table.cell(12, 0).merge(table.cell(12, 7)), "四、文物简介", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(13, 0).merge(table.cell(13, 7)), heritage.description or "（暂无简介）")

    _set_cell_text(table.cell(14, 0).merge(table.cell(14, 7)), "五、照片说明（现场采集照片）", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    _set_cell_text(table.cell(15, 0).merge(table.cell(16, 1)), "照片说明", bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER)
    photo_cell = table.cell(15, 2).merge(table.cell(16, 7))
    _insert_photos_into_cell(photo_cell, list(heritage.photos.all()))

    document.add_paragraph("")
    p_coord_title = document.add_paragraph("区块2坐标点信息")
    p_coord_title.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    for run in p_coord_title.runs:
        run.bold = True
        run.font.size = Pt(14)
        run.font.name = "宋体"

    for line in _format_coord_point_lines(heritage.coord_list):
        p_coord_line = document.add_paragraph(line)
        p_coord_line.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        for run in p_coord_line.runs:
            run.font.size = Pt(12)
            run.font.name = "宋体"

    document.add_paragraph("")
    sign = document.add_paragraph(
        "采集人：{0}    审核人：{1}    日期：{2}".format(
            heritage.collector.get_full_name() if heritage.collector else "",
            heritage.reviewer.get_full_name() if heritage.reviewer else "",
            timezone.localtime(heritage.collected_at).strftime("%Y-%m-%d") if heritage.collected_at else "",
        )
    )
    sign.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    for run in sign.runs:
        run.font.size = Pt(14)
        run.font.name = "宋体"

    output = io.BytesIO()
    document.save(output)
    output.seek(0)

    filename = f"不可移动文物采集登记表_{_safe_docx_fragment(heritage.survey_code)}_{_safe_docx_fragment(heritage.name)}.docx"
    return output, filename


@login_required
def export_immovable_heritage_docx_view(request, pk):
    """
    一键导出：鄯善县不可移动文物采集登记表（.docx）
    - python-docx 动态绘制复杂表格
    - A4 纵向 + 标准页边距
    - 单元格合并 + 现场照片插入
    """
    from .models import ImmovableHeritage

    heritage = get_object_or_404(
        ImmovableHeritage.objects.select_related("collector", "input_by", "reviewer").prefetch_related("photos"),
        pk=pk,
    )
    output, filename = _build_immovable_heritage_docx_stream(heritage)
    response = FileResponse(
        output,
        as_attachment=True,
        filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    response["Content-Type"] = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return response

